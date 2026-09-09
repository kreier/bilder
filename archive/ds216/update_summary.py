# unique edition for my network drive

import os
import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# 🔧 CONFIGURATION
# Path to investigate (network drive UNC path in Windows)
TARGET_PATH = r"\\ds216\photo"

# Output files will be stored in the same location as the Python script
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_FILE = os.path.join(SCRIPT_DIR, "pictures_summary.csv")
EXCEL_FILE = os.path.join(SCRIPT_DIR, "pictures_summary.xlsx")

def get_folder_info(base_path):
    folder_data = []
    for item in os.listdir(base_path):
        item_path = os.path.join(base_path, item)
        if os.path.isdir(item_path):
            folder_size = 0
            num_pictures = 0
            num_other_files = 0
            num_subfolders = 0

            for root, dirs, files in os.walk(item_path):
                num_subfolders += len(dirs)
                for file in files:
                    file_path = os.path.join(root, file)
                    try:
                        folder_size += os.path.getsize(file_path)
                    except OSError:
                        continue
                    
                    if file.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tiff', '.webp')):
                        num_pictures += 1
                    else:
                        num_other_files += 1

            size_gb = round(folder_size / (1024**3), 2)
            folder_data.append([item, folder_size, size_gb, num_pictures, num_other_files, num_subfolders])
    return folder_data

def write_to_csv_and_excel(data, csv_file=CSV_FILE, excel_file=EXCEL_FILE):
    headers = ["Folder Name", "Size (bytes)", "Size (GB)", "Number of Pictures", "Number of Other Files", "Number of Subfolders"]
    
    # Write CSV
    with open(csv_file, mode="w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(data)
    
    # Write Excel
    df = pd.DataFrame(data, columns=headers)
    df.to_excel(excel_file, index=False)

    # Format Excel with openpyxl
    wb = load_workbook(excel_file)
    ws = wb.active

    # Bold headers
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[col_letter].width = adjusted_width

    # Format numeric columns
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):  # Size (bytes) and Size (GB)
        for cell in row:
            if cell.column == 2:  # bytes
                cell.number_format = '#,##0'
            elif cell.column == 3:  # GB
                cell.number_format = '0.00'

    wb.save(excel_file)

if __name__ == "__main__":
    folder_info = get_folder_info(TARGET_PATH)
    write_to_csv_and_excel(folder_info)
    print(f"Reports generated at script location:\n- {CSV_FILE}\n- {EXCEL_FILE}")
